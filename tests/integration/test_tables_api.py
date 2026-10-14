"""Модуль 3 через API: загрузка, предложение, предпросмотр, шаблон маппинга, пакет задач."""

import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from gosdoc.core.config import get_settings
from gosdoc.core.db import StoredFile
from gosdoc.core.queue import get_enqueue
from gosdoc.core.storage import get_storage
from gosdoc.tables.api import get_llm
from gosdoc.tables.service import run_table_job
from openpyxl import load_workbook
from sqlalchemy.orm import Session, sessionmaker

from tests.conftest import MakeUser, login
from tests.tables_spec import EXPECTED, FIXTURES, SPEC
from tests.unit.test_tables_mapping import fake_llm

pytestmark = pytest.mark.integration


@pytest.fixture
def tables_client(client: TestClient, sessions: sessionmaker[Session]) -> TestClient:
    app = client.app
    storage = app.dependency_overrides[get_storage]()  # type: ignore[attr-defined]

    def run_now(queue: str, actor: str, job_id: int) -> None:
        assert (queue, actor) == ("tables", "process_table_job")
        run_table_job(sessions, storage, get_settings(), job_id)

    app.dependency_overrides[get_enqueue] = lambda: run_now  # type: ignore[attr-defined]
    app.dependency_overrides[get_llm] = lambda: fake_llm("{}", [])  # type: ignore[attr-defined]
    return client


def upload(client: TestClient, path: Path) -> int:
    r = client.post("/api/files", files={"file": (path.name, path.read_bytes())})
    assert r.status_code == 201, r.text
    return int(r.json()["id"])


def test_transfer_a_to_b_via_api(
    tables_client: TestClient, make_user: MakeUser, sessions: sessionmaker[Session]
) -> None:
    c = tables_client
    login(c, make_user().login)
    xlsx, csv = upload(c, FIXTURES / "source.xlsx"), upload(c, FIXTURES / "source.csv")
    template = upload(c, FIXTURES / "template.xlsx")

    r = c.post("/api/tables/suggest", json={"source_file_id": xlsx, "template_file_id": template})
    assert r.status_code == 200, r.text
    suggestion = r.json()
    assert suggestion["is_draft"] and suggestion["target_columns"][0] == "ФИО"
    assert {c["target"]: c["op"] for c in suggestion["spec"]["columns"]}["ФИО"] == "fio_join"

    spec = SPEC.model_dump()
    r = c.post("/api/tables/preview", json={"source_file_id": xlsx, "spec": spec})
    assert r.json()["summary"] == EXPECTED and r.json()["rows"][0]["Пол"] == "Мужской"

    r = c.post(
        "/api/tables/mappings",
        json={"name": "Реестр выплат", "spec": spec, "template_file_id": template},
    )
    assert r.status_code == 201, r.text
    mapping = r.json()
    with sessions() as db:  # шаблон хранится без строк данных и без срока удаления
        kept = db.get(StoredFile, mapping["target_template_file_id"])
        assert kept is not None and kept.expires_at is None and not kept.contains_pdn

    r = c.post(
        "/api/tables/jobs", json={"source_file_ids": [xlsx, csv], "mapping_id": mapping["id"]}
    )
    assert r.status_code == 201, r.text
    for job in r.json():
        done = c.get(f"/api/tables/jobs/{job['id']}").json()
        assert done["status"] == "done", done
        assert {k: done["result"][k] for k in EXPECTED} == EXPECTED
        content = c.get(f"/api/files/{done['result']['result_file_id']}/content").content
        ws = load_workbook(io.BytesIO(content)).active
        assert (
            ws.max_row == 1 + EXPECTED["rows_written"]
            and ws.cell(2, 1).value == "Иванов Иван Иванович"
        )
        report = c.get(f"/api/files/{done['result']['report_file_id']}/content")
        assert report.status_code == 200


def test_bad_mapping_fails_job_with_reason(tables_client: TestClient, make_user: MakeUser) -> None:
    c = tables_client
    login(c, make_user().login)
    source, template = upload(c, FIXTURES / "source.csv"), upload(c, FIXTURES / "template.xlsx")
    spec = {"columns": [{"target": "ФИО", "source": "Нет такой"}]}
    r = c.post("/api/tables/preview", json={"source_file_id": source, "spec": spec})
    assert r.status_code == 400 and "Нет такой" in r.json()["detail"]
    r = c.post(
        "/api/tables/jobs",
        json={"source_file_ids": [source], "spec": spec, "template_file_id": template},
    )
    job = c.get(f"/api/tables/jobs/{r.json()[0]['id']}").json()
    assert job["status"] == "failed" and "Нет такой" in job["error"]


def test_mappings_and_jobs_are_private(tables_client: TestClient, make_user: MakeUser) -> None:
    c = tables_client
    login(c, make_user().login)
    template = upload(c, FIXTURES / "template.xlsx")
    body = {"name": "Личный", "spec": SPEC.model_dump(), "template_file_id": template}
    private = c.post("/api/tables/mappings", json=body).json()
    shared = c.post("/api/tables/mappings", json={**body, "name": "Общий", "shared": True}).json()

    login(c, make_user().login)  # коллега из той же организации
    names = {m["name"] for m in c.get("/api/tables/mappings").json()}
    assert "Общий" in names and private["id"] not in {
        m["id"] for m in c.get("/api/tables/mappings").json()
    }
    assert c.delete(f"/api/tables/mappings/{shared['id']}").status_code == 403

    login(c, make_user(org_id=2).login)  # другая организация не видит даже общий
    assert shared["id"] not in {m["id"] for m in c.get("/api/tables/mappings").json()}
    r = c.post("/api/tables/jobs", json={"source_file_ids": [template], "mapping_id": shared["id"]})
    assert r.status_code == 404
