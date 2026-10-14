"""Модуль 2 через API: черновик, версии, проверки, утверждение, DOCX/PDF, бланк."""

import io
import json
import shutil
from pathlib import Path

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from gosdoc.core.config import Settings, get_settings
from gosdoc.core.llm import LLMClient
from gosdoc.core.queue import get_enqueue
from gosdoc.core.storage import get_storage
from gosdoc.letters.api import get_llm
from gosdoc.letters.service import run_pdf_job
from sqlalchemy.orm import Session, sessionmaker

from tests.conftest import MakeUser, login

pytestmark = pytest.mark.integration
FIXTURES = Path(__file__).parent.parent / "fixtures" / "letters"
DRAFT = {
    "title": "О ремонте дороги по ул. Садовой",
    "greeting": "Уважаемый Пётр Петрович!",
    "paragraphs": ["Ваше обращение рассмотрено.", "Ремонт будет выполнен в ближайшее время."],
    "closing": "",
}


@pytest.fixture
def prompts() -> list[str]:
    return []


@pytest.fixture
def letters_client(
    client: TestClient, sessions: sessionmaker[Session], prompts: list[str]
) -> TestClient:
    app = client.app
    storage = app.dependency_overrides[get_storage]()  # type: ignore[attr-defined]

    def handler(request: httpx.Request) -> httpx.Response:
        prompts.append(json.loads(request.content)["messages"][1]["content"])
        content = "Черновик:\n" + json.dumps(DRAFT, ensure_ascii=False)
        return httpx.Response(200, json={"choices": [{"message": {"content": content}}]})

    def run_now(queue: str, actor: str, job_id: int) -> None:
        assert (queue, actor) == ("letters", "process_letter_pdf")
        run_pdf_job(sessions, storage, get_settings(), job_id)

    llm = LLMClient(Settings(), transport=httpx.MockTransport(handler))
    app.dependency_overrides[get_llm] = lambda: llm  # type: ignore[attr-defined]
    app.dependency_overrides[get_enqueue] = lambda: run_now  # type: ignore[attr-defined]
    return client


def upload(c: TestClient, path: Path) -> int:
    r = c.post("/api/files", files={"file": (path.name, path.read_bytes())})
    assert r.status_code == 201, r.text
    return int(r.json()["id"])


def new_letter(c: TestClient, **extra: object) -> dict:  # type: ignore[type-arg]
    body = {
        "type": "citizen_reply",
        "description": "Сообщить, что ремонт запланирован на 2027 год",
        "addressee": "Петрову П.П.\nул. Садовая, д. 1, кв. 2,\nг. Энск, 123456",
        "requisites": {"incoming_number": "П-45", "incoming_date": "2026-10-12"},
        **extra,
    }
    r = c.post("/api/letters", json=body)
    assert r.status_code == 201, r.text
    return r.json()  # type: ignore[no-any-return]


def test_citizen_reply_lifecycle(
    letters_client: TestClient, make_user: MakeUser, prompts: list[str]
) -> None:
    c = letters_client
    admin = make_user(roles=["sys_admin", "user"])
    author = make_user(department="Отдел ЖКХ")
    manager = make_user(roles=["user", "manager"], department="Отдел ЖКХ")

    login(c, admin.login)
    settings = {
        "blank": {"org_name": "АДМИНИСТРАЦИЯ ГОРОДА ЭНСКА", "details": ["ул. Ленина, д. 1"]},
        "outgoing_number_pattern": r"\d{2}-\d{2}/\d+",
        "signer_position": "Глава администрации",
        "signer_name": "И.И. Иванов",
    }
    assert c.put("/api/admin/organization", json=settings).status_code == 200
    assert (
        c.put("/api/admin/organization", json={"outgoing_number_pattern": "("}).status_code == 422
    )

    login(c, author.login)
    letter = new_letter(c, incoming_file_id=upload(c, FIXTURES / "appeal.docx"))
    assert "дорога по ул. Садовой" in prompts[0]  # текст обращения дошёл до модели
    assert letter["is_draft"] and letter["deadline"] == "2026-11-11"
    v1 = letter["versions"][0]
    assert v1["author_id"] is None and v1["meta"]["prompt_version"] == "citizen_reply-v1"
    assert v1["body"]["title"] == DRAFT["title"]
    assert any("конкретный срок" in s["message"] for s in letter["checks"]["style"])

    edit = {
        "body": {
            **DRAFT,
            "paragraphs": ["Ваше обращение рассмотрено.", "Ремонт включён в план на 2027 год."],
        },
        "addressee": letter["addressee"],
        "executor": "Сидорова А.С.",
        "outgoing_number": "01-15",
        "date": "2026-10-20",
        "requisites": {
            "incoming_number": "П-45",
            "incoming_date": "2026-10-12",
            "executor_phone": "45-67-90",
        },
    }
    r = c.put(f"/api/letters/{letter['id']}", json=edit)
    assert r.status_code == 200, r.text
    assert [v["version"] for v in r.json()["versions"]] == [1, 2]
    assert r.json()["checks"]["number_ok"] is False and r.json()["checks"]["style"] == []

    assert c.post(f"/api/letters/{letter['id']}/approve").status_code == 403  # не руководитель
    login(c, manager.login)
    assert c.post(f"/api/letters/{letter['id']}/approve").status_code == 400  # номер не по маске
    login(c, author.login)
    assert (
        c.put(
            f"/api/letters/{letter['id']}", json={**edit, "outgoing_number": "01-15/123"}
        ).status_code
        == 200
    )

    docx = c.post(f"/api/letters/{letter['id']}/docx").json()
    content = c.get(f"/api/files/{docx['docx_file_id']}/content").content
    doc = Document(io.BytesIO(content))
    header = doc.tables[0].rows[0].cells[0].text
    assert "АДМИНИСТРАЦИЯ ГОРОДА ЭНСКА" in header and "20.10.2026 № 01-15/123" in header
    assert "Ремонт включён в план на 2027 год." in [p.text for p in doc.paragraphs]
    assert doc.tables[1].rows[0].cells[1].text == "И.И. Иванов"

    login(c, manager.login)
    r = c.post(f"/api/letters/{letter['id']}/approve")
    assert r.status_code == 200 and r.json()["is_draft"] is False
    assert c.put(f"/api/letters/{letter['id']}", json=edit).status_code == 409

    login(c, make_user(roles=["user", "manager"], department="Другой отдел").login)
    assert c.get(f"/api/letters/{letter['id']}").status_code == 404


def test_blank_upload_and_render(letters_client: TestClient, make_user: MakeUser) -> None:
    c = letters_client
    admin = make_user(roles=["sys_admin", "user"])
    login(c, admin.login)
    file_id = upload(c, FIXTURES / "blank.docx")
    r = c.post("/api/letter-blanks", json={"name": "Бланк письма", "template_file_id": file_id})
    assert r.status_code == 201, r.text
    blank = r.json()
    assert "paragraphs" in blank["fields"]["variables"]

    letter = new_letter(c, type="cover", blank_id=blank["id"])
    docx = c.post(f"/api/letters/{letter['id']}/docx").json()
    text = "\n".join(
        p.text
        for p in Document(
            io.BytesIO(c.get(f"/api/files/{docx['docx_file_id']}/content").content)
        ).paragraphs
    )
    assert "Ваше обращение рассмотрено." in text and "{{" not in text

    login(c, make_user().login)  # загружать бланки может только администратор системы
    own = upload(c, FIXTURES / "blank.docx")
    assert (
        c.post("/api/letter-blanks", json={"name": "x", "template_file_id": own}).status_code == 403
    )


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice не установлен")
def test_pdf_via_worker(letters_client: TestClient, make_user: MakeUser) -> None:
    c = letters_client
    login(c, make_user().login)
    letter = new_letter(c, type="request")
    r = c.post(f"/api/letters/{letter['id']}/pdf")
    assert r.status_code == 202, r.text
    job = c.get(f"/api/letters/jobs/{r.json()['job_id']}").json()
    assert job["status"] == "done", job
    pdf = c.get(f"/api/files/{job['result']['pdf_file_id']}/content").content
    assert pdf.startswith(b"%PDF")
