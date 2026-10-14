"""Оформление по ГОСТ Р 7.0.97-2016 проверяется по XML документа."""

import io
import zipfile
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from gosdoc.letters.gost import LetterDoc, blank_variables, build_docx, render_blank

from tests.tables_spec import FIXTURES

LETTER = LetterDoc(
    org_name="АДМИНИСТРАЦИЯ ГОРОДА ЭНСКА",
    org_details=["ул. Ленина, д. 1, г. Энск, 123456", "Тел. (123) 45-67-89"],
    date=date(2026, 10, 20),
    number="01-15/123",
    incoming_number="П-45",
    incoming_date=date(2026, 10, 12),
    addressee=["Петрову П.П.", "ул. Садовая, д. 1, кв. 2,", "г. Энск, 123456"],
    title="О ремонте дороги по ул. Садовой",
    greeting="Уважаемый Пётр Петрович!",
    paragraphs=["Ваше обращение рассмотрено.", "Ремонт включён в план на 2027 год."],
    attachments=["план ремонта на 2 л. в 1 экз."],
    signer_position="Глава администрации",
    signer_name="И.И. Иванов",
    executor="Сидорова Анна Сергеевна",
    executor_phone="(123) 45-67-90",
)


def test_page_font_and_layout() -> None:
    d = Document(io.BytesIO(build_docx(LETTER)))
    s = d.sections[0]
    sizes = (
        s.page_width,
        s.page_height,
        s.left_margin,
        s.right_margin,
        s.top_margin,
        s.bottom_margin,
    )
    assert [round(v.mm) for v in sizes] == [210, 297, 20, 10, 20, 20]  # в DOCX хранятся twips
    assert d.styles["Normal"].font.name == "Times New Roman"

    texts = [p.text for p in d.paragraphs]
    body = d.paragraphs[texts.index("Ваше обращение рассмотрено.")]
    assert body.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert round(body.paragraph_format.first_line_indent.cm, 2) == 1.25
    assert body.runs[0].font.size == Pt(14)
    greeting = d.paragraphs[texts.index("Уважаемый Пётр Петрович!")]
    assert greeting.alignment == WD_ALIGN_PARAGRAPH.CENTER
    # Порядок реквизитов: заголовок → обращение → текст → приложение → подпись → исполнитель.
    order = [
        "О ремонте дороги по ул. Садовой",
        "Уважаемый Пётр Петрович!",
        "Ваше обращение рассмотрено.",
        "Приложение: план ремонта на 2 л. в 1 экз.",
        "Сидорова Анна Сергеевна",
    ]
    assert [texts.index(t) for t in order] == sorted(texts.index(t) for t in order)
    assert d.paragraphs[texts.index("Сидорова Анна Сергеевна")].runs[0].font.size == Pt(10)

    header, sign = d.tables
    left, right = (c.text for c in header.rows[0].cells)
    assert left.startswith("АДМИНИСТРАЦИЯ ГОРОДА ЭНСКА")
    assert "20.10.2026 № 01-15/123" in left and "На № П-45 от 12.10.2026" in left
    assert right.startswith("Петрову П.П.")
    position, name = sign.rows[0].cells
    assert position.text == "Глава администрации" and name.text == "И.И. Иванов"
    assert name.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT

    xml = zipfile.ZipFile(io.BytesIO(build_docx(LETTER))).read("word/document.xml").decode()
    assert 'w:left="1134"' in xml and 'w:right="567"' in xml  # 20 и 10 мм в twips


def test_blank_template() -> None:
    blank = (FIXTURES.parent / "letters" / "blank.docx").read_bytes()
    assert {"org_name", "paragraphs", "signer_name"} <= blank_variables(blank)
    text = "\n".join(p.text for p in Document(io.BytesIO(render_blank(blank, LETTER))).paragraphs)
    assert "АДМИНИСТРАЦИЯ ГОРОДА ЭНСКА" in text and "20.10.2026 № 01-15/123" in text
    assert "Ремонт включён в план на 2027 год." in text and "{{" not in text
