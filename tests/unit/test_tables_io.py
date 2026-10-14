"""Перенос А → Б на фикстурах: без потерь, с отчётом, в оформлении шаблона."""

import io

import pytest
from gosdoc.tables.tableio import (
    TableReadError,
    header_only_template,
    read_columns,
    read_table,
    write_report,
    write_result,
)
from gosdoc.tables.transforms import apply_mapping
from openpyxl import load_workbook

from tests.tables_spec import EXPECTED, FIXTURES, SPEC


@pytest.mark.parametrize("name", ["source.xlsx", "source.csv", "source.ods"])
def test_all_formats_transfer_identically(name: str) -> None:
    df = read_table((FIXTURES / name).read_bytes(), name)
    assert list(df.index[:3]) == [2, 3, 4]  # номера строк как в файле
    result = apply_mapping(SPEC, df)
    assert result.summary() == EXPECTED
    assert result.data["ФИО"].tolist() == [
        "Иванов Иван Иванович",
        "Петрова Анна Сергеевна",
        "Сидоров Пётр",
        "Новиков Сергей Николаевич",
    ]
    assert result.data["Дата рождения"].tolist() == [
        "15.03.1985",
        "01.07.1990",
        "02.11.1979",
        "23.02.1980",
    ]
    assert result.data["Пол"].tolist() == ["Мужской", "Женский", "Мужской", "Мужской"]
    assert result.data["Сумма"].tolist() == ["12500.00", "8300.50", "15000.00", "11200.00"]


def test_result_written_into_template_and_report_lists_every_lost_row() -> None:
    df = read_table((FIXTURES / "source.xlsx").read_bytes(), "source.xlsx")
    result = apply_mapping(SPEC, df)
    template = (FIXTURES / "template.xlsx").read_bytes()

    out = load_workbook(io.BytesIO(write_result(result, template, "template.xlsx"))).active
    assert out.title == "Реестр" and out.cell(1, 1).font.bold
    assert out.max_row == 1 + 4  # строка-пример шаблона заменена данными
    assert out.cell(2, 1).value == "Иванов Иван Иванович"
    assert out.cell(2, 6).value == 12500.0 and out.cell(5, 6).number_format == "#,##0.00"

    report = load_workbook(io.BytesIO(write_report(result, list(df.columns))))
    rows = list(report["Не перенесено"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == EXPECTED["rows_filtered"] + EXPECTED["rows_rejected"]
    by_row = {r[0]: r for r in rows}
    assert by_row[5][1] == "отклонена" and "ИНН: неверный ИНН" in by_row[5][2]
    assert "Пол: значения «не указан» нет в справочнике" in by_row[7][2]
    assert by_row[8][1] == "отфильтрована" and by_row[8][3] == "Волков"
    assert "СНИЛС: неверный СНИЛС" in by_row[12][2]
    assert "ИНН: обязательное поле не заполнено" in by_row[13][2]


def test_template_helpers() -> None:
    template = (FIXTURES / "template.xlsx").read_bytes()
    assert read_columns(template, "template.xlsx")[0] == "ФИО"
    clean = load_workbook(io.BytesIO(header_only_template(template, "template.xlsx"))).active
    assert clean.max_row == 1  # строка-пример с данными не хранится
    with pytest.raises(TableReadError, match="не поддерживается"):
        read_table(b"x", "data.txt")
    with pytest.raises(TableReadError, match="заголовки"):
        read_table(b";a\n1;2\n", "x.csv")
