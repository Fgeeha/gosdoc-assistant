"""Пересборка фикстур модуля 3: uv run python tests/fixtures/tables/build_fixtures.py

Люди вымышленные. ИНН и СНИЛС составлены с верными контрольными суммами,
кроме намеренно испорченных строк.
"""

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).parent
SOURCE_COLUMNS = [
    "Фамилия",
    "Имя",
    "Отчество",
    "ДР",
    "Пол",
    "ИНН",
    "СНИЛС",
    "Сумма, руб.",
    "Статус",
    "Отдел",
]
TEMPLATE_COLUMNS = ["ФИО", "Дата рождения", "Пол", "ИНН", "СНИЛС", "Сумма", "Подразделение"]


def inn12(base: str) -> str:
    n = [int(c) for c in base]
    c11 = sum(w * x for w, x in zip((7, 2, 4, 10, 3, 5, 9, 4, 6, 8), n, strict=True)) % 11 % 10
    n.append(c11)
    c12 = sum(w * x for w, x in zip((3, 7, 2, 4, 10, 3, 5, 9, 4, 6, 8), n, strict=True)) % 11 % 10
    return base + f"{c11}{c12}"


def snils(base: str) -> str:
    control = sum(int(c) * (9 - i) for i, c in enumerate(base)) % 101
    control = 0 if control == 100 else control
    return f"{base[:3]}-{base[3:6]}-{base[6:]} {control:02d}"


def broken(value: str) -> str:
    """Меняет последнюю цифру — контрольная сумма перестаёт сходиться."""
    return value[:-1] + str((int(value[-1]) + 1) % 10)


ROWS = [
    # Фамилия, Имя, Отчество, ДР, Пол, ИНН, СНИЛС, Сумма, Статус, Отдел
    [
        "Иванов",
        "Иван",
        "Иванович",
        datetime(1985, 3, 15),
        "м",
        inn12("7701234567"),
        snils("123456789"),
        "12 500,00",
        "активен",
        "Отдел кадров",
    ],
    [
        "Петрова",
        "Анна",
        "Сергеевна",
        "1990-07-01",
        "ж",
        inn12("7702345678"),
        snils("234567891"),
        "8 300,50 руб.",
        "активен",
        "Бухгалтерия",
    ],
    [
        "Сидоров",
        "Пётр",
        "",
        "02/11/1979",
        "М",
        inn12("7703456789"),
        snils("345678912"),
        "15000",
        "Активен",
        "Отдел кадров",
    ],
    [
        "Кузнецова",
        "Мария",
        "Олеговна",
        "31.12.2001",
        "ж",
        broken(inn12("7704567891")),
        snils("456789123"),
        "9 999,99",
        "активен",
        "Бухгалтерия",
    ],
    [
        "Смирнов",
        "Алексей",
        "Павлович",
        "32.13.2000",
        "м",
        inn12("7705678912"),
        snils("567891234"),
        "4 000",
        "активен",
        "Отдел кадров",
    ],
    [
        "Попова",
        "Елена",
        "Викторовна",
        "05.05.1995",
        "не указан",
        inn12("7706789123"),
        snils("678912345"),
        "6 100,00",
        "активен",
        "Бухгалтерия",
    ],
    [
        "Волков",
        "Дмитрий",
        "Андреевич",
        "10.10.1988",
        "м",
        inn12("7707891234"),
        snils("789123456"),
        "20 000",
        "уволен",
        "Отдел кадров",
    ],
    [
        "Морозова",
        "Ольга",
        "Игоревна",
        "12.06.1992",
        "ж",
        inn12("7708912345"),
        snils("891234567"),
        "7 450,00",
        "уволен",
        "Бухгалтерия",
    ],
    [
        "Новиков",
        "Сергей",
        "Николаевич",
        "23.02.1980",
        "м",
        inn12("7709123456"),
        snils("912345678"),
        "11 200,00",
        "активен",
        "Бухгалтерия",
    ],
    ["", "", "", "", "", "", "", "", "", ""],  # пустая строка — не данные
    [
        "Фёдорова",
        "Наталья",
        "Юрьевна",
        "08.03.1987",
        "ж",
        inn12("7710234567"),
        broken(snils("102345678")),
        "5 500,00",
        "активен",
        "Отдел кадров",
    ],
    [
        "Егоров",
        "Кирилл",
        "Сергеевич",
        "",
        "м",
        "",
        snils("103456789"),
        "3 300,00",
        "активен",
        "Отдел кадров",
    ],
]


def build() -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(SOURCE_COLUMNS)
    for row in ROWS:
        ws.append(row)
        if isinstance(row[3], datetime):
            ws.cell(ws.max_row, 4).number_format = "DD.MM.YYYY"
    wb.save(HERE / "source.xlsx")

    as_text = [[v.strftime("%d.%m.%Y") if isinstance(v, datetime) else v for v in r] for r in ROWS]
    df = pd.DataFrame(as_text, columns=SOURCE_COLUMNS)
    df.to_csv(HERE / "source.csv", sep=";", index=False, encoding="cp1251")
    df.to_excel(HERE / "source.ods", index=False, engine="odf")

    tpl = Workbook()
    ws = tpl.active
    ws.title = "Реестр"
    ws.append(TEMPLATE_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    # Строка-пример: заменяется данными, её оформление переносится на все строки.
    ws.append(
        [
            "Пример Пример Пример",
            "01.01.2000",
            "Мужской",
            "000000000000",
            "000-000-000 00",
            0,
            "Отдел",
        ]
    )
    ws.cell(2, 6).number_format = "#,##0.00"
    tpl.save(HERE / "template.xlsx")


if __name__ == "__main__":
    build()
