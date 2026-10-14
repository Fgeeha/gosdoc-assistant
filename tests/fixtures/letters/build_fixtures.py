"""Пересборка фикстур модуля 2: uv run python tests/fixtures/letters/build_fixtures.py

Обращение и заявитель вымышленные.
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).parent

APPEAL = [
    "Главе администрации",
    "от Петрова Петра Петровича,",
    "проживающего по адресу: г. Энск, ул. Садовая, д. 1, кв. 2",
    "",
    "Обращение",
    "Прошу сообщить, когда будет отремонтирована дорога по ул. Садовой, "
    "на которой с весны образовались глубокие ямы. Также прошу разъяснить, "
    "включена ли улица в план ремонта на текущий год.",
    "",
    "12.10.2026                         П.П. Петров",
]


def build() -> None:
    doc = Document()
    for line in APPEAL:
        doc.add_paragraph(line)
    doc.save(HERE / "appeal.docx")

    # Бланк организации — шаблон docxtpl с переменными gosdoc.letters.gost.CONTEXT_VARS.
    blank = Document()
    style = blank.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(14)
    for line in (
        "{{ org_name }}",
        "{% for d in org_details %}{{ d }} {% endfor %}",
        "{{ date }} № {{ number }}",
        "{{ incoming_ref }}",
        "{{ addressee }}",
        "{{ title }}",
        "{{ greeting }}",
        "{% for p in paragraphs %}{{ p }}\n{% endfor %}",
        "{{ closing }}",
        "{{ signer_position }}    {{ signer_name }}",
        "{{ executor }} {{ executor_phone }}",
    ):
        blank.add_paragraph(line)
    blank.save(HERE / "blank.docx")


if __name__ == "__main__":
    build()
