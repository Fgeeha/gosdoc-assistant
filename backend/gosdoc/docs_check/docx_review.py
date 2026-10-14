"""Результат проверки в DOCX: примечания Word к фрагментам (ADR-0010) и отчёт-таблица."""

import io
from collections.abc import Sequence
from pathlib import PurePath

from docx import Document
from docx.document import Document as DocumentType
from docx.shared import Pt
from docx.text.paragraph import Paragraph

from gosdoc.docs_check.checker import FindingDraft

AUTHOR = "ГосДок (черновик)"
SEVERITY = {"critical": "Критично", "warning": "Внимание", "info": "К сведению"}


def comment_text(f: FindingDraft) -> str:
    lines = [f"[{SEVERITY[f.severity]}] {f.message}"]
    if f.norm:
        lines.append(f"Норма: {f.norm}")
    if f.quote:
        lines.append(f"Цитата: «{f.quote[:1500]}»")
    if f.recommendation:
        lines.append(f"Предлагаемая правка: {f.recommendation}")
    return "\n".join(lines)


def _anchor(paragraphs: list[Paragraph], fragment: str) -> Paragraph | None:
    probe = fragment[:60].strip()
    for p in paragraphs:
        if probe and probe in " ".join(p.text.split()):
            return p
    return None


def reviewed_docx(
    original: bytes, filename: str, paragraphs: Sequence[str], findings: Sequence[FindingDraft]
) -> bytes:
    """Для DOCX — исходный документ с примечаниями; для других форматов — текст в новом DOCX."""
    if PurePath(filename).suffix.lower() == ".docx":
        doc: DocumentType = Document(io.BytesIO(original))
    else:
        doc = Document()
        for text in paragraphs:
            doc.add_paragraph(text)
    all_paragraphs = list(doc.paragraphs) + [
        p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs
    ]
    for f in findings:
        target = _anchor(all_paragraphs, f.doc_fragment) or doc.paragraphs[0]
        runs = target.runs or [target.add_run("")]
        doc.add_comment(runs, text=comment_text(f), author=AUTHOR, initials="ГД")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def report_docx(
    source_name: str, summary: dict[str, object], findings: Sequence[FindingDraft]
) -> bytes:
    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.add_heading(f"Отчёт о проверке: {source_name}", level=1)
    doc.add_paragraph(
        "Черновик: замечания подготовлены автоматически и требуют проверки сотрудником. "
        "Замечания по содержанию приводятся только с дословной цитатой нормы."
    )
    doc.add_paragraph(
        f"Ссылок на НПА: {summary.get('refs_total', 0)}; замечаний: {len(findings)} "
        f"(критичных {summary.get('critical', 0)}, внимание {summary.get('warning', 0)}, "
        f"к сведению {summary.get('info', 0)})."
    )
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, title in zip(
        table.rows[0].cells,
        ("Критичность", "Фрагмент документа", "Норма", "Замечание", "Рекомендация"),
        strict=True,
    ):
        cell.text = title
    for f in findings:
        row = table.add_row().cells
        row[0].text = SEVERITY[f.severity]
        row[1].text = f.doc_fragment[:400]
        row[2].text = f.norm + (f"\n«{f.quote[:800]}»" if f.quote else "")
        row[3].text = f.message
        row[4].text = f.recommendation
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
