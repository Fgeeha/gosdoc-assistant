"""Извлечение текста из документов: DOCX, ODT, PDF (с OCR для сканов), HTML, TXT.

Абзацы возвращаются списком: проверка документов привязывает замечания к ним.
OCR (Tesseract, rus) есть только в образе воркера: в API для сканов — ошибка.
"""

import io
import re
import shutil
import subprocess
import tempfile
from html.parser import HTMLParser
from pathlib import Path, PurePath
from typing import ClassVar

import pdfplumber
from docx import Document

FORMATS = (".docx", ".odt", ".pdf", ".html", ".htm", ".txt")
OCR_DPI = 300


class TextExtractionError(ValueError):
    pass


class _HtmlText(HTMLParser):
    _BLOCK: ClassVar[set[str]] = {
        "p",
        "div",
        "br",
        "li",
        "tr",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "section",
        "article",
    }
    _SKIP: ClassVar[set[str]] = {"script", "style", "head"}

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self.skip += 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self.skip:
            self.skip -= 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.skip:
            self.parts.append(data)


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1251")


def _docx(data: bytes) -> list[str]:
    doc = Document(io.BytesIO(data))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return [p.text for p in doc.paragraphs] + cells


def _odt(data: bytes) -> list[str]:
    from odf import teletype
    from odf import text as odf_text
    from odf.opendocument import load

    doc = load(io.BytesIO(data))
    return [teletype.extractText(p) for p in doc.getElementsByType(odf_text.P)]


def _ocr(data: bytes) -> list[str]:
    if shutil.which("tesseract") is None:
        raise TextExtractionError(
            "PDF без текстового слоя (скан): распознавание доступно в фоновой проверке"
        )
    import pypdfium2 as pdfium

    lines: list[str] = []
    pdf = pdfium.PdfDocument(data)
    with tempfile.TemporaryDirectory() as tmp:
        for i, page in enumerate(pdf):
            image = page.render(scale=OCR_DPI / 72).to_pil()
            path = Path(tmp) / f"p{i}.png"
            image.save(path)
            out = subprocess.run(  # noqa: S603
                ["tesseract", str(path), "stdout", "-l", "rus"],  # noqa: S607
                check=True,
                capture_output=True,
                timeout=300,
            )
            # Tesseract разделяет абзацы пустой строкой; строки внутри абзаца склеиваются,
            # иначе ссылка «части 1 статьи 12 … № 59-ФЗ» рвётся по переносу строки.
            text = out.stdout.decode("utf-8", "replace")
            lines += [
                re.sub(r"-\n(?=[а-яё])", "", block).replace("\n", " ")
                for block in re.split(r"\n\s*\n", text)
            ]
    return lines


def _pdf(data: bytes, ocr: bool) -> list[str]:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        lines = [line for page in pdf.pages for line in (page.extract_text() or "").splitlines()]
    if any(line.strip() for line in lines):
        return lines
    if not ocr:
        raise TextExtractionError(
            "В PDF нет текстового слоя (скан) — загрузите DOCX или PDF с текстом"
        )
    return _ocr(data)


def extract_paragraphs(data: bytes, filename: str, ocr: bool = False) -> list[str]:
    """Абзацы документа (пустые убраны). ocr=True — распознавать сканы PDF."""
    ext = PurePath(filename).suffix.lower()
    try:
        if ext == ".docx":
            raw = _docx(data)
        elif ext == ".odt":
            raw = _odt(data)
        elif ext == ".pdf":
            raw = _pdf(data, ocr)
        elif ext in (".html", ".htm"):
            parser = _HtmlText()
            parser.feed(_decode(data))
            raw = "".join(parser.parts).split("\n")
        elif ext == ".txt":
            raw = _decode(data).splitlines()
        else:
            raise TextExtractionError(
                f"Формат {ext or 'без расширения'} не поддерживается: {', '.join(FORMATS)}"
            )
    except TextExtractionError:
        raise
    except Exception as e:  # повреждённый файл любого формата
        raise TextExtractionError(f"Не удалось прочитать документ: {type(e).__name__}") from e
    return [re.sub(r"\s+", " ", p).strip() for p in raw if p and p.strip()]


def extract_text(data: bytes, filename: str, ocr: bool = False) -> str:
    return "\n".join(extract_paragraphs(data, filename, ocr))
