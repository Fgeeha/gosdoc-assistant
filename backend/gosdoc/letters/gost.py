"""Оформление письма по ГОСТ Р 7.0.97-2016.

Реквизиты подставляет только этот код; модель пишет лишь смысловые блоки.
Встроенный бланк собирается python-docx; бланк организации, загруженный
администратором, — DOCX-шаблон docxtpl с теми же переменными (см. CONTEXT_VARS).
"""

import datetime as dt
import io
from dataclasses import asdict, dataclass, field

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from docxtpl import DocxTemplate

# Поля страницы (п. 6.1): левое 20 мм, правое 10 мм, верхнее и нижнее 20 мм.
MARGINS_MM = {"left": 20, "right": 10, "top": 20, "bottom": 20}
FONT = "Times New Roman"
TEXT_PT = 14  # п. 6.2: 12–14 пт
REQUISITES_PT = 12
EXECUTOR_PT = 10  # отметка об исполнителе — меньшим размером
INDENT_CM = 1.25  # абзацный отступ

# Переменные, доступные в загруженном бланке docxtpl.
CONTEXT_VARS = (
    "org_name",
    "org_details",
    "date",
    "number",
    "incoming_ref",
    "addressee",
    "title",
    "greeting",
    "paragraphs",
    "closing",
    "attachments",
    "signer_position",
    "signer_name",
    "executor",
    "executor_phone",
)


@dataclass
class LetterDoc:
    org_name: str
    org_details: list[str] = field(default_factory=list)  # адрес, телефон, e-mail, ОГРН/ИНН
    date: dt.date | None = None
    number: str = ""
    incoming_number: str = ""
    incoming_date: dt.date | None = None
    addressee: list[str] = field(default_factory=list)  # строки реквизита «Адресат»
    title: str = ""  # заголовок к тексту: «О …»
    greeting: str = ""  # «Уважаемый Иван Иванович!»
    paragraphs: list[str] = field(default_factory=list)
    closing: str = ""
    attachments: list[str] = field(default_factory=list)
    signer_position: str = ""
    signer_name: str = ""  # «И.О. Фамилия» — инициалы перед фамилией (п. 5.23)
    executor: str = ""
    executor_phone: str = ""

    @property
    def date_text(self) -> str:
        return self.date.strftime("%d.%m.%Y") if self.date else "_____________"

    @property
    def incoming_ref(self) -> str:
        if not (self.incoming_number or self.incoming_date):
            return ""
        d = self.incoming_date.strftime("%d.%m.%Y") if self.incoming_date else "__________"
        return f"На № {self.incoming_number or '_____'} от {d}"

    def context(self) -> dict[str, object]:
        ctx = {k: v for k, v in asdict(self).items() if k in CONTEXT_VARS}
        ctx.update(
            date=self.date_text, number=self.number or "_____", incoming_ref=self.incoming_ref
        )
        ctx["addressee"] = "\n".join(self.addressee)
        return ctx


def _fmt(
    p: Paragraph,
    size: int,
    *,
    align: WD_PARAGRAPH_ALIGNMENT | None = None,
    bold: bool = False,
    indent: bool = False,
    space_after: int = 0,
) -> Paragraph:
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(space_after), 1.0
    if align is not None:
        p.alignment = align
    if indent:
        pf.first_line_indent = Cm(INDENT_CM)
    for run in p.runs:
        run.font.size, run.font.bold = Pt(size), bold
    return p


def _add(container: object, text: str, size: int = TEXT_PT, **kw: object) -> Paragraph:
    p = container.add_paragraph(text)  # type: ignore[attr-defined]
    return _fmt(p, size, **kw)  # type: ignore[arg-type]


def _no_borders(table: object) -> None:
    tbl_pr = table._tbl.tblPr  # type: ignore[attr-defined]
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)


def _cell(cell: _Cell, lines: list[str], size: int, **kw: object) -> None:
    cell.paragraphs[0].text = lines[0] if lines else ""
    _fmt(cell.paragraphs[0], size, **kw)  # type: ignore[arg-type]
    for line in lines[1:]:
        _fmt(cell.add_paragraph(line), size, **kw)  # type: ignore[arg-type]


def build_docx(doc: LetterDoc) -> bytes:
    """Письмо на встроенном бланке (угловое расположение реквизитов бланка)."""
    d = Document()
    section = d.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin, section.right_margin = Mm(MARGINS_MM["left"]), Mm(MARGINS_MM["right"])
    section.top_margin, section.bottom_margin = Mm(MARGINS_MM["top"]), Mm(MARGINS_MM["bottom"])
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(TEXT_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    # Шапка: слева — бланк, дата и номер, ссылка на входящий; справа — адресат.
    header = d.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.LEFT
    _no_borders(header)
    left, right = header.rows[0].cells
    left.width, right.width = Cm(9), Cm(9)
    _cell(left, [doc.org_name], REQUISITES_PT, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    for line in doc.org_details:
        _fmt(left.add_paragraph(line), REQUISITES_PT - 2, align=WD_ALIGN_PARAGRAPH.CENTER)
    _fmt(left.add_paragraph(""), REQUISITES_PT)
    _fmt(left.add_paragraph(f"{doc.date_text} № {doc.number or '_____'}"), REQUISITES_PT)
    if doc.incoming_ref:
        _fmt(left.add_paragraph(doc.incoming_ref), REQUISITES_PT)
    _cell(right, doc.addressee or [""], TEXT_PT)

    _add(d, "", space_after=0)
    if doc.title:
        _add(d, doc.title, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=12)
    if doc.greeting:
        _add(d, doc.greeting, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    for text in doc.paragraphs:
        _add(d, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True)
    if doc.closing:
        _add(d, doc.closing, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True)

    if doc.attachments:
        _add(d, "")
        if len(doc.attachments) == 1:
            _add(d, f"Приложение: {doc.attachments[0]}", align=WD_ALIGN_PARAGRAPH.LEFT)
        else:
            _add(d, "Приложение:", align=WD_ALIGN_PARAGRAPH.LEFT)
            for i, a in enumerate(doc.attachments, 1):
                _add(d, f"{i}. {a}", align=WD_ALIGN_PARAGRAPH.LEFT)

    # Подпись: должность — от левого поля, расшифровка — у правого (п. 5.23).
    _add(d, "")
    _add(d, "")
    sign = d.add_table(rows=1, cols=2)
    _no_borders(sign)
    pos, name = sign.rows[0].cells
    pos.width, name.width = Cm(10), Cm(8)
    _cell(pos, doc.signer_position.splitlines() or [""], TEXT_PT, align=WD_ALIGN_PARAGRAPH.LEFT)
    _cell(name, [doc.signer_name], TEXT_PT, align=WD_ALIGN_PARAGRAPH.RIGHT)

    if doc.executor:  # отметка об исполнителе — внизу слева (п. 5.27)
        _add(d, "")
        _add(d, "")
        _add(d, doc.executor, EXECUTOR_PT, align=WD_ALIGN_PARAGRAPH.LEFT)
        if doc.executor_phone:
            _add(d, doc.executor_phone, EXECUTOR_PT, align=WD_ALIGN_PARAGRAPH.LEFT)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_blank(template: bytes, doc: LetterDoc) -> bytes:
    """Письмо на бланке организации (DOCX-шаблон docxtpl)."""
    tpl = DocxTemplate(io.BytesIO(template))
    tpl.render(doc.context(), autoescape=True)
    buf = io.BytesIO()
    tpl.save(buf)
    return buf.getvalue()


def blank_variables(template: bytes) -> set[str]:
    """Какие переменные использует бланк — для проверки при загрузке."""
    return set(DocxTemplate(io.BytesIO(template)).get_undeclared_template_variables())
