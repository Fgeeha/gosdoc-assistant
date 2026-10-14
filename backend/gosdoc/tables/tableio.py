"""Чтение таблиц (XLSX, XLS, CSV, ODS) и запись результата и отчёта в XLSX."""

import csv
import io
from copy import copy
from decimal import Decimal
from pathlib import PurePath
from typing import Literal

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from gosdoc.tables.transforms import FIRST_DATA_ROW, TransferResult

FORMATS = (".xlsx", ".xls", ".csv", ".ods")
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_ENGINES: dict[str, Literal["openpyxl", "xlrd", "odf"]] = {
    ".xlsx": "openpyxl",
    ".xls": "xlrd",
    ".ods": "odf",
}


class TableReadError(ValueError):
    pass


def _ext(filename: str) -> str:
    ext = PurePath(filename).suffix.lower()
    if ext not in FORMATS:
        raise TableReadError(
            f"Формат {ext or 'без расширения'} не поддерживается: {', '.join(FORMATS)}"
        )
    return ext


def _read_csv(data: bytes, nrows: int | None) -> pd.DataFrame:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1251")  # выгрузки из 1С и старого Excel
    try:
        sep = csv.Sniffer().sniff(text[:4096], delimiters=";,\t").delimiter
    except csv.Error:
        sep = ";"
    return pd.read_csv(io.StringIO(text), sep=sep, dtype=str, keep_default_na=False, nrows=nrows)


def read_table(data: bytes, filename: str, nrows: int | None = None) -> pd.DataFrame:
    """Все значения — строки; индекс — номер строки в файле (заголовок — строка 1).

    Полностью пустые строки данными не считаются и пропускаются.
    """
    ext = _ext(filename)
    try:
        if ext == ".csv":
            df = _read_csv(data, nrows)
        else:
            df = pd.read_excel(
                io.BytesIO(data),
                engine=_ENGINES[ext],
                dtype=str,
                keep_default_na=False,
                nrows=nrows,
            )
    except (ValueError, OSError, KeyError) as e:
        raise TableReadError(f"Не удалось прочитать таблицу: {e}") from e
    df.columns = [str(c).strip() for c in df.columns]
    if any(c.startswith("Unnamed:") or not c for c in df.columns):
        raise TableReadError("В первой строке таблицы должны быть заголовки всех колонок")
    df = df.fillna("").astype(str)
    df.index = df.index + FIRST_DATA_ROW
    return df[df.apply(lambda r: any(v.strip() for v in r), axis=1)] if len(df) else df


def read_columns(data: bytes, filename: str) -> list[str]:
    return list(read_table(data, filename, nrows=0).columns)


def _cell_value(value: str, numeric: bool) -> str | float:
    if numeric and value:
        return float(Decimal(value))
    return value


def write_result(result: TransferResult, template: bytes, template_name: str) -> bytes:
    """Данные пишутся в шаблон XLSX под заголовок с сохранением оформления;
    для шаблонов других форматов создаётся новый XLSX с теми же заголовками."""
    columns = list(result.data.columns)
    styles: dict[int, dict[str, object]] = {}
    if _ext(template_name) == ".xlsx":
        wb = load_workbook(io.BytesIO(template))
        ws = _sheet(wb)
        header = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
        # Оформление первой строки данных шаблона переносится на все строки.
        styles = {col: _style(ws.cell(FIRST_DATA_ROW, col)) for col in header.values()}
        if ws.max_row >= FIRST_DATA_ROW:  # примеры строк в шаблоне заменяются данными
            ws.delete_rows(FIRST_DATA_ROW, ws.max_row)
    else:
        wb = Workbook()
        ws = _sheet(wb)
        ws.append(columns)
        header = {c: i for i, c in enumerate(columns, start=1)}
    for offset, record in enumerate(result.data.itertuples(index=False)):
        for name, value in zip(columns, record, strict=True):
            cell = ws.cell(FIRST_DATA_ROW + offset, header[name])
            cell.value = _cell_value(str(value), name in result.numeric_columns)
            for attr, value in styles.get(header[name], {}).items():
                setattr(cell, attr, copy(value))
    return _save(wb)


def write_report(result: TransferResult, source_columns: list[str]) -> bytes:
    """Сводка и все не перенесённые строки с причиной и исходными значениями."""
    wb = Workbook()
    summary = _sheet(wb)
    summary.title = "Сводка"
    labels = {
        "rows_in": "Строк в источнике",
        "rows_out": "Перенесено",
        "rows_filtered": "Отфильтровано правилами",
        "rows_rejected": "Отклонено с ошибками",
        "rows_written": "Строк в результате (после группировки)",
    }
    for key, value in result.summary().items():
        summary.append([labels[key], value])
    summary.append([])
    summary.append(["Перенесено + отфильтровано + отклонено = строк в источнике"])
    summary.column_dimensions["A"].width = 45

    errors = wb.create_sheet("Не перенесено")
    errors.append(["Строка источника", "Статус", "Причина", *source_columns])
    for cell in errors[1]:
        cell.font = Font(bold=True)
    for issue in result.issues:
        errors.append(
            [
                issue.row,
                issue.status,
                issue.reason,
                *(issue.values.get(c, "") for c in source_columns),
            ]
        )
    errors.column_dimensions["C"].width = 60
    return _save(wb)


def _style(cell: Cell | MergedCell) -> dict[str, object]:
    if isinstance(cell, MergedCell) or not cell.has_style:
        return {}
    attrs = ("font", "border", "fill", "number_format", "alignment", "protection")
    return {a: getattr(cell, a) for a in attrs}


def header_only_template(data: bytes, filename: str) -> bytes:
    """Копия шаблона только с заголовком: так сохранённый маппинг не хранит ПДн."""
    if _ext(filename) != ".xlsx":
        wb = Workbook()
        ws = _sheet(wb)
        ws.append(read_columns(data, filename))
        return _save(wb)
    book = load_workbook(io.BytesIO(data))
    for sheet in book.worksheets:
        if sheet.max_row > 1:
            sheet.delete_rows(FIRST_DATA_ROW, sheet.max_row)
    return _save(book)


def _sheet(wb: Workbook) -> Worksheet:
    ws = wb.active
    if not isinstance(ws, Worksheet):
        raise TableReadError("Первый лист книги — не таблица")
    return ws


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
